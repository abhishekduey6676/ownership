"""Build the portfolio PRD from its canonical Markdown without rendering.

Requires python-docx. Run from any directory with Python 3:
    python docs/portfolio/build_prd.py
The generated document contains the same content, with portable source links.
"""

from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/product/PRD.md"
OUTPUT = ROOT / "docs/portfolio/Ownership-PRD.docx"
BASELINE = "ff78ad9bf92252f1c390cbf43ba168a63c696a94"
REMOTE = "https://github.com/abhishekduey6676/ownership/blob/"
TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def href(target):
    if target.startswith("https://"):
        return target
    path = (SOURCE.parent / target).resolve()
    if not path.is_file():
        raise ValueError(f"Broken source link: {target}")
    relative = path.relative_to(ROOT).as_posix()
    # Newly authored documentation is browsable on its documentation branch.
    ref = "codex/ownership-as-built-prd" if relative.endswith("PRD-roadmap-archive.md") else BASELINE
    return REMOTE + ref + "/" + relative


def inline(paragraph, text):
    for token in TOKEN.split(text):
        if not token:
            continue
        if token.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, target = match.groups()
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(href(target), RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "222222")
            props.append(color)
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            props.append(underline)
            run.append(props)
            value = OxmlElement("w:t")
            value.text = label
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1] if token.startswith("`") else token)
            run.bold = token.startswith("**")
            if token.startswith("`"):
                run.font.name = "Consolas"


def table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError("Table column mismatch")
    widths = {2: [2.0, 4.5], 3: [1.25, 2.15, 3.1]}[count]
    result = doc.add_table(rows=0, cols=count)
    result.autofit = False
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    props = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        item = OxmlElement("w:" + edge)
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            item.set(qn("w:" + key), value)
        borders.append(item)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in [("top", "85"), ("bottom", "85"), ("left", "100"), ("right", "100")]:
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    props.append(margins)
    for index, values in enumerate(rows):
        row = result.add_row()
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if index == 0:
            row_props.append(OxmlElement("w:tblHeader"))
        for cell, width, value in zip(row.cells, widths, values):
            cell.width = Inches(width)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.06
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(11)
                if index == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.35)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12
    doc.styles["Title"].font.size = Pt(26)
    doc.styles["Title"].font.bold = True
    for name, size in [("Heading 1", 17), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    doc.styles["Normal"].paragraph_format.widow_control = True
    header = section.header.paragraphs[0]
    header.add_run("OWNERSHIP  /  PRODUCT REQUIREMENTS").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run("Version 1.0  |  2 October 2026                                      ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Ownership Product Requirements Document"
    doc.core_properties.subject = "As built product specification and portfolio edition"
    doc.core_properties.author = "Ownership"
    doc.core_properties.keywords = "Ownership, PRD, AI, product management"
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("|"):
            group = []
            while index < len(lines) and lines[index].startswith("|"):
                group.append(lines[index])
                index += 1
            table(doc, group)
            continue
        heading = re.match(r"^(#{1,3}) (.+)$", line)
        if heading:
            level, text = heading.groups()
            if not re.fullmatch(r"[A-Za-z0-9 ]+", text):
                raise ValueError(f"Nonplain heading: {text}")
            paragraph = doc.add_paragraph(style={1: "Title", 2: "Heading 1", 3: "Heading 2"}[len(level)])
            inline(paragraph, text)
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            # Literal source numbering ensures each independent list starts at one.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            inline(paragraph, line)
        else:
            inline(doc.add_paragraph(), line)
        index += 1
    doc.save(OUTPUT)
    validate(markdown)


def validate(markdown):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        xml = ET.fromstring(archive.read("word/document.xml"))
        extracted = ["".join(node.itertext()) for node in xml.findall(".//w:t", ns)]
        actual = "".join(extracted)
        expected = []
        for line in markdown.splitlines():
            if not line.strip():
                continue
            if line.startswith("|"):
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    continue
                expected.extend(cells)
            else:
                expected.append(re.sub(r"^(?:#{1,3} |- )", "", line))
        plain = "".join(expected)
        plain = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", plain)
        plain = plain.replace("**", "").replace("`", "")
        normalize = lambda value: re.sub(r"\s+", "", value)
        assert normalize(plain) == normalize(actual), "Source and DOCX text diverged"
        assert len(xml.findall(".//w:tblHeader", ns)) == len(xml.findall(".//w:tbl", ns))
        assert not xml.findall(".//w:trHeight", ns), "Unexpected fixed row height"
    print(f"Created {OUTPUT}")
    print(f"Source words: {len(markdown.split())}; text parity and table checks passed")
    print("Visual rendering intentionally not performed at user request")


if __name__ == "__main__":
    build()
